import argparse
import hashlib
import json
import random
from datetime import datetime, timedelta
from pathlib import Path


import csv

import psycopg


from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DOCUMENTS_DIR = PROJECT_ROOT / "sample-documents"


DEPARTMENTS = {
    "Finance": [
        "Q3 Expense Summary",
        "FY2026 Department Budget",
        "Vendor Invoice Register",
        "Travel Reimbursement Report",
        "Monthly Financial Forecast",
        "Accounts Payable Summary",
        "Capital Expenditure Review",
        "Budget Variance Analysis",
        "Quarterly Cost Report",
        "Procurement Expense Overview",
    ],
    "Human Resources": [
        "Employee Handbook",
        "Benefits Overview",
        "Remote Work Policy",
        "New Hire Checklist",
        "Training Schedule",
        "Performance Review Guidelines",
        "Employee Onboarding Guide",
        "Leave and Attendance Policy",
        "Professional Development Plan",
        "Workplace Conduct Guide",
    ],
    "Sales": [
        "Q3 Sales Performance",
        "Product Pricing Sheet",
        "Regional Sales Summary",
        "Lead Pipeline Report",
        "Customer Presentation",
        "Monthly Revenue Report",
        "Sales Territory Plan",
        "Account Growth Analysis",
        "Quarterly Forecast",
        "Sales Enablement Guide",
    ],
    "Operations": [
        "Inventory Status Report",
        "Warehouse Procedure",
        "Vendor Contact List",
        "Equipment Maintenance Schedule",
        "Standard Operating Procedure",
        "Facilities Inspection Report",
        "Supply Chain Summary",
        "Shipping Process Guide",
        "Inventory Reconciliation Report",
        "Operations KPI Review",
    ],
    "Information Technology": [
        "Security Awareness Guide",
        "System Inventory",
        "Incident Response Checklist",
        "Cloud Access Procedure",
        "Software License Register",
        "Network Maintenance Schedule",
        "Backup and Recovery Procedure",
        "Access Review Report",
        "Patch Management Guide",
        "IT Service Continuity Plan",
    ],
}


UPLOADERS = {
    "Finance": [
        "Alicia Morgan",
        "Daniel Brooks",
        "Priya Shah",
    ],
    "Human Resources": [
        "Monica Reed",
        "Kevin Foster",
        "Sofia Ramirez",
    ],
    "Sales": [
        "Jordan Lee",
        "Marcus Hill",
        "Emily Carter",
    ],
    "Operations": [
        "Nathan Price",
        "Olivia Bennett",
        "Carlos Mendes",
    ],
    "Information Technology": [
        "Ethan Walker",
        "Maya Patel",
        "Liam Chen",
    ],
}


FILE_EXTENSIONS = [
    "pdf",
    "docx",
    "xlsx",
    "csv",
    "txt",
    "png",
]


DESCRIPTION_TEMPLATES = [
    "Internal business document for the {department} department.",
    "Operational reference material used by the {department} team.",
    "Departmental document supporting routine {department} activities.",
    "Internal report prepared for the {department} department.",
    "Business record maintained by the {department} team.",
]


def build_business_content(title, department, uploader, uploaded_at):
    return (
        f"{title}\n\n"
        f"Department: {department}\n"
        f"Document Owner: {uploader}\n"
        f"Document Date: {uploaded_at.strftime('%Y-%m-%d')}\n\n"
        f"Purpose\n"
        f"This document contains synthetic business information created for "
        f"the {department} department as part of the Cloud Document Management "
        f"System migration project.\n\n"
        f"Operational Summary\n"
        f"The information in this document represents realistic internal "
        f"business content used to test document storage, metadata management, "
        f"cross-cloud migration, and data-integrity validation.\n\n"
        f"Migration Classification\n"
        f"Source Platform: Amazon Web Services (AWS)\n"
        f"Target Platform: Microsoft Azure\n"
        f"Data Classification: Synthetic / Non-sensitive\n"
    )


def create_txt_file(path, title, department, uploader, uploaded_at):
    content = build_business_content(
        title, department, uploader, uploaded_at
    )
    path.write_text(content, encoding="utf-8")


def create_csv_file(path, title, department, uploader, uploaded_at):
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)

        writer.writerow(
            ["Record ID", "Department", "Metric", "Value", "Owner", "Date"]
        )

        for row_number in range(1, 11):
            writer.writerow(
                [
                    row_number,
                    department,
                    f"{title} Metric {row_number}",
                    random.randint(100, 5000),
                    uploader,
                    uploaded_at.strftime("%Y-%m-%d"),
                ]
            )


def create_docx_file(path, title, department, uploader, uploaded_at):
    document = Document()

    document.add_heading(title, level=0)
    document.add_paragraph(f"Department: {department}")
    document.add_paragraph(f"Document Owner: {uploader}")
    document.add_paragraph(
        f"Document Date: {uploaded_at.strftime('%Y-%m-%d')}"
    )

    document.add_heading("Purpose", level=1)
    document.add_paragraph(
        f"This synthetic document represents internal business material "
        f"for the {department} department."
    )

    document.add_heading("Operational Summary", level=1)
    document.add_paragraph(
        "This file is part of a controlled cross-cloud migration dataset "
        "used to test document storage, metadata migration, application "
        "functionality, and integrity validation."
    )

    document.add_heading("Migration Information", level=1)
    document.add_paragraph("Source Platform: Amazon Web Services (AWS)")
    document.add_paragraph("Target Platform: Microsoft Azure")
    document.add_paragraph("Data Classification: Synthetic / Non-sensitive")

    document.save(path)


def create_xlsx_file(path, title, department, uploader, uploaded_at):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Business Data"

    worksheet.append(
        ["Record ID", "Department", "Metric", "Value", "Owner", "Date"]
    )

    for row_number in range(1, 16):
        worksheet.append(
            [
                row_number,
                department,
                f"{title} Metric {row_number}",
                random.randint(100, 10000),
                uploader,
                uploaded_at.strftime("%Y-%m-%d"),
            ]
        )

    metadata_sheet = workbook.create_sheet("Metadata")
    metadata_sheet.append(["Field", "Value"])
    metadata_sheet.append(["Document Title", title])
    metadata_sheet.append(["Department", department])
    metadata_sheet.append(["Document Owner", uploader])
    metadata_sheet.append(
        ["Document Date", uploaded_at.strftime("%Y-%m-%d")]
    )
    metadata_sheet.append(["Source Cloud", "AWS"])
    metadata_sheet.append(["Target Cloud", "Azure"])
    metadata_sheet.append(["Classification", "Synthetic / Non-sensitive"])

    workbook.save(path)


def create_pdf_file(path, title, department, uploader, uploaded_at):
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter

    pdf.setTitle(title)

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(72, height - 72, title)

    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, height - 105, f"Department: {department}")
    pdf.drawString(72, height - 125, f"Document Owner: {uploader}")
    pdf.drawString(
        72,
        height - 145,
        f"Document Date: {uploaded_at.strftime('%Y-%m-%d')}",
    )

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(72, height - 185, "Purpose")

    pdf.setFont("Helvetica", 10)
    pdf.drawString(
        72,
        height - 205,
        "Synthetic business document for cross-cloud migration testing.",
    )

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(72, height - 245, "Migration Information")

    pdf.setFont("Helvetica", 10)
    pdf.drawString(72, height - 265, "Source Platform: Amazon Web Services")
    pdf.drawString(72, height - 285, "Target Platform: Microsoft Azure")
    pdf.drawString(
        72,
        height - 305,
        "Classification: Synthetic / Non-sensitive",
    )

    pdf.save()


def create_png_file(path, title, department, uploader, uploaded_at):
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)

    draw.text((60, 60), title, fill="black")
    draw.text((60, 120), f"Department: {department}", fill="black")
    draw.text((60, 160), f"Document Owner: {uploader}", fill="black")
    draw.text(
        (60, 200),
        f"Document Date: {uploaded_at.strftime('%Y-%m-%d')}",
        fill="black",
    )

    draw.text(
        (60, 280),
        "Cloud Document Management System",
        fill="black",
    )
    draw.text(
        (60, 330),
        "Source: Amazon Web Services (AWS)",
        fill="black",
    )
    draw.text(
        (60, 370),
        "Target: Microsoft Azure",
        fill="black",
    )
    draw.text(
        (60, 410),
        "Classification: Synthetic / Non-sensitive",
        fill="black",
    )

    image.save(path)


def slugify(value):
    """Convert a document title into a storage-safe filename."""
    return (
        value.lower()
        .replace("&", "and")
        .replace(" ", "-")
        .replace("/", "-")
    )


def generate_upload_date(index):
    """Generate predictable historical upload dates for migration testing."""
    base_date = datetime(2026, 1, 15, 9, 0, 0)
    return base_date + timedelta(days=index * 4, hours=index % 8)



def build_document_records():
    """Build realistic metadata records and physical files for the dataset."""
    random.seed(42)

    records = []
    record_number = 1

    for department, titles in DEPARTMENTS.items():
        for title in titles:
            extension = random.choice(FILE_EXTENSIONS)
            uploader = random.choice(UPLOADERS[department])

            filename = f"{slugify(title)}.{extension}"

            department_folder = (
                department.lower()
                .replace(" ", "-")
            )

            object_key = f"{department_folder}/{filename}"

            description = random.choice(
                DESCRIPTION_TEMPLATES
            ).format(department=department)

            uploaded_at = generate_upload_date(record_number - 1)

            department_path = SAMPLE_DOCUMENTS_DIR / department_folder
            department_path.mkdir(parents=True, exist_ok=True)

            file_path = department_path / filename

            file_generators = {
                "txt": create_txt_file,
                "csv": create_csv_file,
                "docx": create_docx_file,
                "xlsx": create_xlsx_file,
                "pdf": create_pdf_file,
                "png": create_png_file,
            }

            generator = file_generators[extension]

            generator(
                file_path,
                title,
                department,
                uploader,
                uploaded_at,
            )

            file_size = file_path.stat().st_size

            with file_path.open("rb") as file:
                checksum = hashlib.sha256(file.read()).hexdigest()

            record = {
                "record_number": record_number,
                "title": title,
                "original_filename": filename,
                "object_key": object_key,
                "description": description,
                "department": department,
                "uploaded_by": uploader,
                "uploaded_at": uploaded_at.isoformat(),
                "file_type": extension,
                "file_size": file_size,
                "storage_provider": "local",
                "storage_location": str(
                    department_path.relative_to(PROJECT_ROOT)
                ),
                "checksum": checksum,
            }

            records.append(record)
            record_number += 1

    return records




def build_manifest_from_existing_files():
    """Build metadata from the existing generated files without recreating them."""
    random.seed(42)

    records = []
    record_number = 1

    for department, titles in DEPARTMENTS.items():
        for title in titles:
            uploader = random.choice(UPLOADERS[department])

            department_folder = (
            department.lower()
            .replace(" ", "-")
            )

            department_path = SAMPLE_DOCUMENTS_DIR / department_folder

            file_stem = slugify(title)

            matching_files = list(department_path.glob(f"{file_stem}.*"))

            if len(matching_files) != 1:
                raise FileNotFoundError(
                    f"Expected exactly one source file for '{title}', "
                    f"found {len(matching_files)}."
                )

            file_path = matching_files[0]
            filename = file_path.name
            extension = file_path.suffix.lstrip(".")

            object_key = f"{department_folder}/{filename}"

            description = random.choice(
                DESCRIPTION_TEMPLATES
            ).format(department=department)

            uploaded_at = generate_upload_date(record_number - 1)

            file_size = file_path.stat().st_size
            with file_path.open("rb") as file:
                     checksum = hashlib.sha256(file.read()).hexdigest()
           

            with file_path.open("rb") as file:
                checksum = hashlib.sha256(file.read()).hexdigest()

            record = {
                "record_number": record_number,
                "title": title,
                "original_filename": filename,
                "object_key": object_key,
                "description": description,
                "department": department,
                "uploaded_by": uploader,
                "uploaded_at": uploaded_at.isoformat(),
                "file_type": extension,
                "file_size": file_size,
                "storage_provider": "local",
                "storage_location": str(
                    department_path.relative_to(PROJECT_ROOT)
                ),
                "checksum": checksum,
            }

            records.append(record)
            record_number += 1

    return records

def seed_postgres_from_manifest():
    manifest_path = SAMPLE_DOCUMENTS_DIR / "manifest.json"

    with manifest_path.open("r", encoding="utf-8") as file:
        records = json.load(file)

    connection = psycopg.connect(
        "dbname=aws_azure_document_migration"
    )

    insert_sql = """
        INSERT INTO documents (
            record_number,
            original_filename,
            object_key,
            title,
            description,
            department,
            uploaded_by,
            uploaded_at,
            file_type,
            file_size,
            storage_provider,
            storage_location,
            checksum
        )
        VALUES (
            %(record_number)s,
            %(original_filename)s,
            %(object_key)s,
            %(title)s,
            %(description)s,
            %(department)s,
            %(uploaded_by)s,
            %(uploaded_at)s,
            %(file_type)s,
            %(file_size)s,
            %(storage_provider)s,
            %(storage_location)s,
            %(checksum)s
        )
        ON CONFLICT (record_number)
        DO UPDATE SET
            original_filename = EXCLUDED.original_filename,
            object_key = EXCLUDED.object_key,
            title = EXCLUDED.title,
            description = EXCLUDED.description,
            department = EXCLUDED.department,
            uploaded_by = EXCLUDED.uploaded_by,
            uploaded_at = EXCLUDED.uploaded_at,
            file_type = EXCLUDED.file_type,
            file_size = EXCLUDED.file_size,
            storage_provider = EXCLUDED.storage_provider,
            storage_location = EXCLUDED.storage_location,
            checksum = EXCLUDED.checksum,
            updated_at = CURRENT_TIMESTAMP;
    """

    try:
        with connection.cursor() as cursor:
            for record in records:
                cursor.execute(
                    insert_sql,
                    record,
                )

        connection.commit()

        print(
            f"Seeded PostgreSQL records: {len(records)}"
        )

    except Exception:
        connection.rollback()
        raise

    finally:
        connection.close()


def write_manifest_from_existing_files():
    records = build_manifest_from_existing_files()

    manifest_path = SAMPLE_DOCUMENTS_DIR / "manifest.json"

    with manifest_path.open("w", encoding="utf-8") as file:
        json.dump(records, file, indent=2)

    print(f"Generated records: {len(records)}")
    print(f"Manifest written to: {manifest_path}")
    print("")
    print("First record:")
    print(json.dumps(records[0], indent=2))
    print("")
    print("Last record:")
    print(json.dumps(records[-1], indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Generate or seed document migration test data."
    )

    parser.add_argument(
        "--seed-db",
        action="store_true",
        help="Load the existing manifest into PostgreSQL.",
    )

    args = parser.parse_args()

    if args.seed_db:
        seed_postgres_from_manifest()
    else:
        write_manifest_from_existing_files()